"""Three analyst-style layouts of the same logical model (model.py).

  S1  "statements"   one sheet per statement (Inputs, IS, BS, CF, DCF), $ in millions,
                     standard labels, fiscal-quarter headers
  S2  "one sheet"    everything stacked on one sheet (Model), $ in thousands, the filer's own
                     statement labels, period-end dates as headers
  S3  "driver model" Drivers, Quarterly and Annual sheets plus Valuation, $ in millions,
                     abbreviated labels, and two analyst hardcodes (a Q4 gross margin override
                     and a typed other-income estimate) an update must leave alone

Every workbook also has an empty Scratch sheet. A style maps each logical cell to a sheet and a
coordinate, scales amounts to its units, and writes labels and headers. `render` turns formula
templates into cell references.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from filing_update_evals import model as M

STANDARD = {
    "is": {
        "revenue": "Revenue",
        "cogs": "Cost of revenue",
        "gross_profit": "Gross profit",
        "sga": "Selling, general and administrative",
        "other_opex": "Other operating expenses, net",
        "operating_income": "Operating income",
        "interest_expense": "Interest expense, net",
        "other_nonop": "Other non-operating income (expense), net",
        "pretax_income": "Income before income taxes",
        "income_tax": "Income tax expense",
        "other_after_tax": "Other items after tax, incl. non-controlling interests",
        "net_income": "Net income attributable to the company",
        "da": "Memo: depreciation and amortization (cash flow statement)",
        "ebitda": "EBITDA (operating income + D&A)",
    },
    "bs": {
        "cash": "Cash and cash equivalents",
        "receivables": "Accounts receivable, net",
        "inventory": "Inventories",
        "other_current_assets": "Other current assets",
        "current_assets": "Total current assets",
        "ppe": "Property and equipment, net",
        "goodwill": "Goodwill",
        "intangibles": "Intangible assets, net",
        "other_noncurrent_assets": "Other non-current assets",
        "total_assets": "Total assets",
        "payables": "Accounts payable",
        "short_term_debt": "Short-term borrowings",
        "current_ltd": "Current portion of long-term debt",
        "other_current_liabilities": "Other current liabilities",
        "current_liabilities": "Total current liabilities",
        "long_term_debt": "Long-term debt",
        "other_noncurrent_liabilities": "Other non-current liabilities",
        "total_liabilities": "Total liabilities",
        "redeemable_nci": "Redeemable non-controlling interests",
        "equity": "Shareholders' equity (company)",
        "nci": "Non-controlling interests",
        "total_equity": "Total equity",
        "total_liabilities_equity": "Total liabilities and equity",
        "check": "Check: total assets - total liabilities and equity (should be 0)",
    },
    "cf": {
        "net_income": "Net income",
        "da": "Depreciation and amortization",
        "wc_other": "Working capital and other operating items",
        "cfo": "Cash from operating activities",
        "capex": "Capital expenditures",
        "other_investing": "Other investing activities",
        "cfi": "Cash from investing activities",
        "cff": "Cash from financing activities",
        "fx_other": "Exchange rate effect, restricted cash and other",
        "net_change": "Net change in cash",
        "begin_cash": "Cash at beginning of period",
        "end_cash": "Cash at end of period",
        "fcf": "Free cash flow (operating cash flow + capital expenditures)",
    },
    "ytd": {
        "cfo": "Cash from operating activities, year to date (as reported)",
        "capex": "Capital expenditures, year to date (as reported)",
        "cfi": "Cash from investing activities, year to date (as reported)",
        "cff": "Cash from financing activities, year to date (as reported)",
        "da": "Depreciation and amortization, year to date (as reported)",
    },
    "val": {
        "quarter": "Fiscal quarter",
        "ebit": "EBIT",
        "tax_on_ebit": "Tax on EBIT",
        "nopat": "NOPAT",
        "da": "Depreciation and amortization",
        "capex": "Capital expenditures",
        "wc": "Working capital and other",
        "ufcf": "Unlevered free cash flow",
        "period": "Discount period (years from the balance sheet date)",
        "discount": "Discount factor",
        "pv": "PV of unlevered free cash flow",
        "ltm_revenue": "Revenue, last four reported quarters",
        "ltm_ebitda": "EBITDA, last four reported quarters",
        "stub_years": "Rest of the current fiscal year (years)",
        "stub_ufcf": "Unlevered free cash flow, rest of the current fiscal year",
        "stub_discount": "Discount factor, rest of the current fiscal year",
        "pv_stub": "PV, rest of the current fiscal year",
        "terminal_value": "Terminal value (Gordon growth, end of last forecast year)",
        "pv_terminal": "PV of terminal value",
        "enterprise_value": "Enterprise value",
        "cash": "Less: cash and cash equivalents (latest reported balance sheet)",
        "short_term_debt": "Short-term borrowings (latest reported balance sheet)",
        "current_ltd": "Current portion of long-term debt (latest reported balance sheet)",
        "long_term_debt": "Long-term debt (latest reported balance sheet)",
        "net_debt": "Net debt",
        "nci": "Non-controlling interests, incl. redeemable (latest reported balance sheet)",
        "equity_value": "Equity value",
        "shares": "Shares outstanding",
        "per_share": "Equity value per share ($)",
        "ev_ltm_ebitda": "Enterprise value / EBITDA, last four reported quarters (x)",
    },
}

ABBREVIATED = {
    "is": {
        "revenue": "Rev",
        "cogs": "COGS",
        "gross_profit": "GP",
        "sga": "SG&A",
        "other_opex": "Other opex",
        "operating_income": "EBIT",
        "interest_expense": "Int exp, net",
        "other_nonop": "Other non-op",
        "pretax_income": "EBT",
        "income_tax": "Tax",
        "other_after_tax": "Other after-tax / NCI",
        "net_income": "NI",
        "da": "D&A (CF stmt)",
        "ebitda": "EBITDA",
    },
    "bs": {
        "cash": "Cash",
        "receivables": "AR",
        "inventory": "Inventory",
        "other_current_assets": "Other CA",
        "current_assets": "Total CA",
        "ppe": "PP&E",
        "goodwill": "Goodwill",
        "intangibles": "Intangibles",
        "other_noncurrent_assets": "Other NCA",
        "total_assets": "Total assets",
        "payables": "AP",
        "short_term_debt": "ST debt",
        "current_ltd": "Current LTD",
        "other_current_liabilities": "Other CL",
        "current_liabilities": "Total CL",
        "long_term_debt": "LT debt",
        "other_noncurrent_liabilities": "Other NCL",
        "total_liabilities": "Total liabilities",
        "redeemable_nci": "Redeemable NCI",
        "equity": "Equity",
        "nci": "NCI",
        "total_equity": "Total equity",
        "total_liabilities_equity": "Total L&E",
        "check": "Check (A - L - E)",
    },
    "cf": {
        "net_income": "NI",
        "da": "D&A",
        "wc_other": "WC & other",
        "cfo": "CFO",
        "capex": "Capex",
        "other_investing": "Other CFI",
        "cfi": "CFI",
        "cff": "CFF",
        "fx_other": "FX & other",
        "net_change": "Net change in cash",
        "begin_cash": "Begin cash",
        "end_cash": "End cash",
        "fcf": "FCF",
    },
    "ytd": {
        "cfo": "CFO YTD (reported)",
        "capex": "Capex YTD (reported)",
        "cfi": "CFI YTD (reported)",
        "cff": "CFF YTD (reported)",
        "da": "D&A YTD (reported)",
    },
}

INPUT_LABELS = {
    "last_q": "Last reported quarter of FY{fy}",
    "shares": "Shares outstanding, latest 10-Q/10-K cover ({shares_unit})",
    "growth_fy": "Revenue growth, FY{fy} quarters vs. same quarter of FY{prev}",
    "growth": "Revenue growth by forecast year",
    "gross_margin": "Gross margin",
    "sga_pct": "SG&A, % of revenue",
    "other_opex_pct": "Other operating expenses, % of revenue",
    "da_pct": "Depreciation and amortization, % of revenue",
    "capex_pct": "Capital expenditures, % of revenue",
    "tax_rate": "Tax rate",
    "dso": "Days sales outstanding (on revenue)",
    "dio": "Days inventory outstanding (on cost of revenue)",
    "dpo": "Days payables outstanding (on cost of revenue)",
    "payout": "Shareholder payouts and other financing, % of net income",
    "interest": "Interest expense, net, per year ({money_unit})",
    "other_nonop_q": "Other non-operating income, per quarter ({money_unit})",
    "other_after_tax_q": "Other items after tax, per quarter ({money_unit})",
    "other_investing_q": "Other investing cash flow, per quarter ({money_unit})",
    "fx_other_q": "Exchange rate effect and other cash flow, per quarter ({money_unit})",
    "wacc": "WACC",
    "terminal_growth": "Terminal growth rate",
    "q4_gross_margin": "Q4 gross margin (analyst override)",
}
INPUT_ORDER = (
    "last_q",
    "shares",
    None,
    "growth_fy",
    "growth",
    "gross_margin",
    "sga_pct",
    "other_opex_pct",
    "da_pct",
    "capex_pct",
    "tax_rate",
    "dso",
    "dio",
    "dpo",
    "payout",
    "interest",
    "other_nonop_q",
    "other_after_tax_q",
    "other_investing_q",
    "fx_other_q",
    None,
    "wacc",
    "terminal_growth",
)

SCRATCH = "Scratch"
GRID = dict(zip(M.COLUMNS, (get_column_letter(i) for i in range(3, 3 + len(M.COLUMNS))), strict=True))
"""The 14-column grid of S1 and S2: C..P."""


@dataclass
class Layout:
    """Where every logical cell goes, for one style and one filer."""

    style: str
    sheets: tuple[str, ...]
    money_scale: float
    shares_scale: float
    money_unit: str
    shares_unit: str
    place: dict[tuple[str, str, str], tuple[str, str]] = field(default_factory=dict)
    """(section, line, col) -> (sheet, coordinate)."""
    static: dict[tuple[str, str], object] = field(default_factory=dict)
    """(sheet, coordinate) -> label, header or status content."""
    header_cells: dict[str, list[tuple[str, str]]] = field(default_factory=dict)
    """Logical column -> its header cells (fiscal label, date, status) on every sheet."""


def _fiscal_labels(data: dict, short: bool) -> dict[str, str]:
    fy = int(data["fy"])

    def name(year: int) -> str:
        return f"FY{year % 100:02d}" if short else f"FY{year}"

    out = {f"P{j}": f"Q{j} {name(fy - 1)}" for j in range(1, 5)}
    out |= {"PY": name(fy - 1), "CY": name(fy)}
    out |= {f"C{j}": f"Q{j} {name(fy)}" for j in range(1, 5)}
    out |= {f"F{n}": f"{name(fy + n)}E" for n in range(1, 5)}
    return out


def _dates(data: dict) -> dict[str, str]:
    import datetime as dt

    def fmt(text: str) -> str:
        return dt.date.fromisoformat(text).strftime("%b %d, %Y")

    cal = data["calendar"]
    out = {f"P{j}": fmt(d) for j, d in enumerate(cal["prev_quarter_ends"], start=1)}
    out["PY"] = fmt(cal["prev_quarter_ends"][3])
    out |= {f"C{j}": fmt(d) for j, d in enumerate(cal["quarter_ends"], start=1)}
    return out


def _labels(style: str, data: dict) -> dict[str, dict[str, str]]:
    if style == "S3":
        out = {sec: dict(STANDARD[sec]) | ABBREVIATED.get(sec, {}) for sec in STANDARD}
        return out
    out = {sec: dict(lines) for sec, lines in STANDARD.items()}
    if style == "S2":  # the filer's own words wherever a line maps to one statement line
        filer = data["labels"]
        for line, label in filer.items():
            if not label:
                continue
            if line in M.IS_TYPED:
                out["is"][line] = label
            elif line in M.BS_TYPED:
                out["bs"][line] = label
            elif line in M.YTD_LINES:
                out["ytd"][line] = f"{label}, year to date (as reported)"
    return out


def layout(style: str, data: dict) -> Layout:
    if style == "S1":
        return _layout_s1(data)
    if style == "S2":
        return _layout_s2(data)
    if style == "S3":
        return _layout_s3(data)
    raise ValueError(style)


def _input_label(name: str, data: dict, lay: Layout) -> str:
    return INPUT_LABELS[name].format(
        fy=data["fy"], prev=int(data["fy"]) - 1, shares_unit=lay.shares_unit, money_unit=lay.money_unit
    )


def _status(col: str, last_ref: str) -> str:
    if col in M.C:
        return f'=IF({col[1]}<={last_ref},"A","E")'
    return {"PY": "A", "CY": "", **{c: "E" for c in M.F}}.get(col, "A")


def _headers(
    lay: Layout, sheet: str, row: int, cols: dict[str, str], labels: dict, dates: dict | None, last_ref: str
) -> None:
    for col, letter in cols.items():
        cells = [(sheet, f"{letter}{row}")]
        lay.static[(sheet, f"{letter}{row}")] = labels[col]
        if dates is not None:
            lay.static[(sheet, f"{letter}{row + 1}")] = dates.get(col, "")
            cells.append((sheet, f"{letter}{row + 1}"))
        status_row = row + (2 if dates is not None else 1)
        lay.static[(sheet, f"{letter}{status_row}")] = _status(col, last_ref)
        cells.append((sheet, f"{letter}{status_row}"))
        lay.header_cells.setdefault(col, []).extend(cells)


def _section(
    lay: Layout, sheet: str, row: int, section: str, lines, cols: dict[str, str], labels: dict
) -> int:
    for line in lines:
        lay.static[(sheet, f"A{row}")] = labels[section][line]
        for col, letter in cols.items():
            lay.place[(section, line, col)] = (sheet, f"{letter}{row}")
        row += 1
    return row


def _inputs(lay: Layout, sheet: str, row: int, data: dict, grow_cols: dict[str, str], plugs: bool) -> int:
    order = list(INPUT_ORDER) + (["q4_gross_margin"] if plugs else [])
    for name in order:
        if name is None:
            row += 1
            continue
        lay.static[(sheet, f"A{row}")] = _input_label(name, data, lay)
        if name == "growth":
            for col, letter in grow_cols.items():
                lay.place[("inp", name, col)] = (sheet, f"{letter}{row}")
        else:
            lay.place[("inp", name, "V")] = (sheet, f"C{row}")
        row += 1
    return row


def _val_scalars(lay: Layout, sheet: str, row: int, labels: dict) -> int:
    for line in M.VAL_LINES:
        if line in ("period", "discount", "pv"):
            continue
        lay.static[(sheet, f"A{row}")] = labels["val"][line]
        lay.place[("val", line, "V")] = (sheet, f"C{row}")
        row += 1
    return row


def _layout_s1(data: dict) -> Layout:
    lay = Layout("S1", ("Inputs", "IS", "BS", "CF", "DCF", SCRATCH), 1e6, 1e6, "$ in millions", "millions")
    labels = _labels("S1", data)
    fiscal = _fiscal_labels(data, short=False)
    last_ref = "Inputs!$C$4"
    name = data["name"]
    # Inputs: scalars in C, per-year growth on the grid's forecast columns.
    lay.static[("Inputs", "A1")] = f"{name}: model inputs"
    lay.static[("Inputs", "A2")] = "$ in millions unless noted; rates as decimals"
    for col in M.F:
        lay.static[("Inputs", f"{GRID[col]}3")] = fiscal[col]
    _inputs(lay, "Inputs", 4, data, {c: GRID[c] for c in M.F}, plugs=False)
    titles = {
        "IS": "income statement",
        "BS": "balance sheet",
        "CF": "cash flow statement",
        "DCF": "unlevered DCF and valuation",
    }
    for sheet, title in titles.items():
        lay.static[(sheet, "A1")] = f"{name}: {title}"
        lay.static[(sheet, "A2")] = "$ in millions"
    is_cols = {c: GRID[c] for c in M.COLUMNS}
    bs_cols = {c: GRID[c] for c in M.BS_COLUMNS}
    cf_cols = {c: GRID[c] for c in M.CF_COLUMNS}
    for sheet, cols in (("IS", is_cols), ("BS", bs_cols), ("CF", cf_cols)):
        _headers(lay, sheet, 3, cols, fiscal, None, last_ref)
    _section(lay, "IS", 6, "is", M.IS_LINES, is_cols, labels)
    _section(lay, "BS", 6, "bs", M.BS_LINES, bs_cols, labels)
    row = _section(lay, "CF", 6, "cf", M.CF_LINES, cf_cols, labels) + 1
    lay.static[("CF", f"A{row}")] = "As reported in each filing (year to date)"
    ytd_cols = {c: GRID[c] for c in (*M.P[:3], "PY", *M.C)}
    _section(lay, "CF", row + 1, "ytd", M.YTD_LINES, ytd_cols, labels)
    q_cols = {c: GRID[c] for c in (*M.C, *M.F)}
    _headers(lay, "DCF", 3, q_cols, fiscal, None, last_ref)
    row = _section(lay, "DCF", 6, "val", M.VAL_Q_LINES, q_cols, labels)
    row = (
        _section(lay, "DCF", row, "val", ("period", "discount", "pv"), {c: GRID[c] for c in M.F}, labels) + 1
    )
    _val_scalars(lay, "DCF", row, labels)
    return lay


def _layout_s2(data: dict) -> Layout:
    lay = Layout("S2", ("Model", SCRATCH), 1e3, 1e3, "$ in thousands", "thousands")
    labels = _labels("S2", data)
    fiscal = _fiscal_labels(data, short=True)
    dates = _dates(data)
    sheet = "Model"
    lay.static[(sheet, "A1")] = f"{data['name']}: three-statement model and DCF"
    lay.static[(sheet, "A2")] = "$ in thousands"
    row = 7
    lay.static[(sheet, f"A{row}")] = "INPUTS"
    first_input = row + 1
    last_ref = f"$C${first_input}"
    _headers(lay, sheet, 3, {c: GRID[c] for c in M.COLUMNS}, fiscal, dates, last_ref)
    row = _inputs(lay, sheet, first_input, data, {c: GRID[c] for c in M.F}, plugs=False) + 1
    for title, section, lines, cols in (
        ("INCOME STATEMENT", "is", M.IS_LINES, M.COLUMNS),
        ("BALANCE SHEET", "bs", M.BS_LINES, M.BS_COLUMNS),
        ("CASH FLOW STATEMENT", "cf", M.CF_LINES, M.CF_COLUMNS),
        ("AS REPORTED IN EACH FILING (YEAR TO DATE)", "ytd", M.YTD_LINES, (*M.P[:3], "PY", *M.C)),
        ("VALUATION", "val", M.VAL_Q_LINES, (*M.C, *M.F)),
    ):
        lay.static[(sheet, f"A{row}")] = title
        row = _section(lay, sheet, row + 1, section, lines, {c: GRID[c] for c in cols}, labels)
        row += 1 if section != "val" else 0
    row = (
        _section(lay, sheet, row, "val", ("period", "discount", "pv"), {c: GRID[c] for c in M.F}, labels) + 1
    )
    _val_scalars(lay, sheet, row, labels)
    return lay


def _layout_s3(data: dict) -> Layout:
    lay = Layout("S3", ("Drivers", "Quarterly", "Annual", "Valuation", SCRATCH), 1e6, 1e6, "$m", "m")
    labels = _labels("S3", data)
    fiscal = _fiscal_labels(data, short=True)
    quarterly = {c: get_column_letter(i) for i, c in enumerate((*M.P, *M.C), start=3)}
    annual = {c: get_column_letter(i) for i, c in enumerate(("PY", "CY", *M.F), start=3)}
    last_ref = "Drivers!$C$4"
    lay.static[("Drivers", "A1")] = f"{data['ticker']} drivers"
    lay.static[("Drivers", "A2")] = "$m; rates as decimals"
    for col in M.F:
        lay.static[("Drivers", f"{annual[col]}3")] = fiscal[col]
    _inputs(lay, "Drivers", 4, data, {c: annual[c] for c in M.F}, plugs=True)
    for sheet, cols in (("Quarterly", quarterly), ("Annual", annual)):
        lay.static[(sheet, "A1")] = f"{data['ticker']} {sheet.lower()} model ($m)"
        _headers(lay, sheet, 2, cols, fiscal, None, last_ref)
        row = 5
        for title, section, lines, ok in (
            ("P&L", "is", M.IS_LINES, M.COLUMNS),
            ("Balance sheet", "bs", M.BS_LINES, M.BS_COLUMNS),
            ("Cash flow", "cf", M.CF_LINES, M.CF_COLUMNS),
            ("Reported YTD", "ytd", M.YTD_LINES, (*M.P[:3], "PY", *M.C)),
        ):
            lay.static[(sheet, f"A{row}")] = title
            row = (
                _section(
                    lay,
                    sheet,
                    row + 1,
                    section,
                    lines,
                    {c: letter for c, letter in cols.items() if c in ok},
                    labels,
                )
                + 1
            )
    lay.static[("Valuation", "A1")] = f"{data['ticker']} DCF ($m)"
    val_q = {c: get_column_letter(i) for i, c in enumerate(M.C, start=3)}
    val_f = {c: get_column_letter(i) for i, c in enumerate(M.F, start=8)}
    _headers(lay, "Valuation", 2, val_q | val_f, fiscal, None, last_ref)
    row = _section(lay, "Valuation", 5, "val", M.VAL_Q_LINES, val_q | val_f, labels)
    row = _section(lay, "Valuation", row, "val", ("period", "discount", "pv"), val_f, labels) + 1
    _val_scalars(lay, "Valuation", row, labels)
    return lay


_PLACEHOLDER = re.compile(r"\{(?:(\w+)\.)?(\w+)(?:@(\w+)(?::(\w+))?)?\}")


def _ref(sheet: str, coord: str, here: str, absolute: bool = False) -> str:
    if absolute:
        letters = "".join(ch for ch in coord if ch.isalpha())
        digits = coord[len(letters) :]
        coord = f"${letters}${digits}"
    return coord if sheet == here else f"{sheet}!{coord}"


def render(template: str, section: str, lay: Layout, here: str) -> str:
    """A formula template -> a spreadsheet formula on sheet `here`."""

    def replace(match: re.Match) -> str:
        sec, line, col, col2 = match.groups()
        sec = sec or section
        if sec == "inp":
            col = col or "V"
            sheet, coord = lay.place[("inp", line, col)]
            return _ref(sheet, coord, here, absolute=(col == "V"))
        col = col or "V"
        try:
            sheet, coord = lay.place[(sec, line, col)]
        except KeyError as exc:
            raise KeyError(f"{lay.style}: no place for {sec}.{line}@{col} (in {template!r})") from exc
        if col2 is None:
            return _ref(sheet, coord, here)
        sheet2, coord2 = lay.place[(sec, line, col2)]
        if sheet2 != sheet:
            raise ValueError(f"{lay.style}: range {sec}.{line}@{col}:{col2} spans sheets")
        return _ref(sheet, f"{coord}:{coord2}", here)

    return "=" + _PLACEHOLDER.sub(replace, template)


def scale(cell: M.Cell, lay: Layout) -> float:
    value = float(cell.content)
    if cell.unit == M.AMOUNT:
        return value / lay.money_scale
    if cell.unit == M.SHARES:
        return value / lay.shares_scale
    return value
