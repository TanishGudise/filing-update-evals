"""Accept a workbook saved by an agent harness, or reject it before anything evaluates it.

An agent may write the .xlsx with any tool, so the file is untrusted. Before openpyxl parses it:
size, entry-count and compression-ratio limits, and no XML part may declare a DOCTYPE or an
entity (the parser would otherwise expand them). Parts that make LibreOffice reach outside the
workbook (macros, external links, data connections, query tables, embedded objects) reject it.

After parsing: every template sheet must be present, no defined names (a name can hide a
function call that no cell shows) except, for a bring-your-own task, the firm workbook's own names
with their own definitions, no array or data-table formulas, and every formula in every cell passes
the same function allowlist as the edits format (plus, for a bring-your-own task, the functions the
firm's updated workbook already uses). Conditional formats and data
validations can't change a value, so they are dropped rather than evaluated.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidationList
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from filing_update_evals.edits import ALLOWED_FUNCTIONS, function_names, normalize_formula

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ENTRIES = 2000
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_RATIO = 200

_FORBIDDEN_PARTS = (
    "xl/vbaproject.bin",
    "xl/externallinks/",
    "xl/connections.xml",
    "xl/querytables/",
    "xl/activex/",
    "xl/embeddings/",
    "xl/ctrlprops/",
)
_DTD = re.compile(rb"<!\s*(DOCTYPE|ENTITY)", re.IGNORECASE)
OUTSIDE_WORKBOOK = re.compile(r"\[|file:|://", re.IGNORECASE)


class Rejected(ValueError):
    pass


def formula_problem(formula: str, extra: frozenset[str] = frozenset()) -> str | None:
    """Why a formula can't be graded, or None. Shared with the edits path. `extra`: functions a
    bring-your-own task allows because the firm's own workbook uses them."""
    names = function_names(formula)
    if names is None:
        return f"formula can't be parsed: {formula[:60]}"
    for name in names:
        if name not in ALLOWED_FUNCTIONS and name not in extra:
            return f"function not allowed: {name}"
    if OUTSIDE_WORKBOOK.search(formula):
        return "reference to another file or a URL"
    return None


def defined_names(wb: Workbook) -> dict[str, str]:
    """Every defined name and what it refers to; a sheet's own names as `Sheet!name`."""
    out = {name: d.attr_text for name, d in wb.defined_names.items()}
    for ws in wb.worksheets:
        out |= {f"{ws.title}!{name}": d.attr_text for name, d in ws.defined_names.items()}
    return out


def _check_zip(data: bytes) -> None:
    if len(data) > MAX_FILE_BYTES:
        raise Rejected(f"file larger than {MAX_FILE_BYTES} bytes")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise Rejected("not an xlsx file") from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ENTRIES:
            raise Rejected("too many parts")
        if sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES:
            raise Rejected("uncompressed size too large")
        for entry in entries:
            name = entry.filename.lower()
            if any(name.startswith(part) for part in _FORBIDDEN_PARTS):
                raise Rejected(f"forbidden part: {entry.filename}")
            if entry.compress_size and entry.file_size / entry.compress_size > MAX_RATIO:
                raise Rejected(f"suspicious compression ratio: {entry.filename}")
            if name.endswith((".xml", ".rels", ".vml")) and _DTD.search(archive.read(entry)):
                raise Rejected(f"DTD or entity declaration in {entry.filename}")


def load_submission(
    source: Path | bytes,
    template_sheets: tuple[str, ...],
    *,
    extra_functions: frozenset[str] = frozenset(),
    allowed_names: dict[str, list[str]] | None = None,
) -> Workbook:
    data = source if isinstance(source, bytes) else Path(source).read_bytes()
    _check_zip(data)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=False, keep_links=False)
    except Exception as exc:  # openpyxl raises many types on malformed files
        raise Rejected(f"openpyxl can't read the file: {type(exc).__name__}") from exc

    missing = [name for name in template_sheets if name not in wb.sheetnames]
    if missing:
        raise Rejected(f"sheet missing: {missing[0]}")
    names = defined_names(wb)
    for name, text in names.items():
        if text not in (allowed_names or {}).get(name, ()):
            raise Rejected(f"defined names are not allowed: {name} = {text}")
    for ws in wb.worksheets:
        ws.conditional_formatting = ConditionalFormattingList()
        ws.data_validations = DataValidationList()
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, ArrayFormula | DataTableFormula):
                    raise Rejected(
                        f"{ws.title}!{cell.coordinate}: array and data-table formulas are not allowed"
                    )
                if isinstance(value, str) and value.startswith("="):
                    problem = formula_problem(value, extra_functions)
                    if problem:
                        raise Rejected(f"{ws.title}!{cell.coordinate}: {problem}")
                    # Same normalization as the edits path (case, `_xlfn.` prefixes).
                    cell.value = normalize_formula(value)
    return wb
